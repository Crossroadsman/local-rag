import os

# filetype handlers
import fitz # `pymupdf` when installing from conda
import docx  # `python-docx` when installing from conda
import openpyxl
import markdown2

from bs4 import BeautifulSoup
import re

# choose whichever file handling library you want for a particular filetype
# import it above then define a handler function below

# handler functions
def extract_txt(file_path):
    with open(file_path, 'r', encoding='utf-8') as file:
        return file.read()

def extract_markdown(file_path):
    with open(file_path, 'r', encoding='utf-8') as file:
        # Convert markdown to HTML
        html_content = markdown2.markdown(file.read())

        # Use BS to parse and remove HTML tags
        bs = BeautifulSoup(html_content, 'html.parser')
        plain_text = bs.get_text()

        return plain_text

def extract_pdf(file_path):
    file = fitz.open(file_path)
    text = ""
    for page in file:
        text += page.get_text("text")

    text = clean_text(text)

    return text

def extract_docx(file_path):
    file = docx.Document(file_path)
    return "\n".join([para.text for para in file.paragraphs])

def extract_xlsx(file_path):
    file = openpyxl.load_workbook(file_path)
    text = []
    for sheet in file.worksheets:
        for row in sheet.iter_rows():
            text.append(" ".join([str(cell.value) for cell in row if cell.value]))
    return "\n".join(text)


# Mapping table
# associate each supported file extension with a handler function
file_type_handlers = {
    ".txt": extract_txt,
    ".md": extract_markdown,
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx
}

# Helper functions
def clean_text(text):
    """This is especially useful for PDFs which can have a lot of non-human-readable text"""

    # Remove non-printable characters
    text = re.sub(r'[^\x20-\x7E\n]', '', text)

    # Remove excessive whitespace or other artifacts
    text = re.sub(r'\s+', ' ', text).strip()

    return text

# Master function
def extract_content(file_path):
    ext = os.path.splitext(file_path)[1]
    if ext in file_type_handlers:
        return file_type_handlers[ext](file_path)
    else:
        raise ValueError(f"No handler available for file type: {ext}")
